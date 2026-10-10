"""数据进出：问卷星文件导入、Excel 导出、导入模板下载。"""
import io
import uuid
from pathlib import Path

from flask import (Blueprint, abort, flash, redirect, render_template, request,
                   send_file, url_for)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import importer
import models
from auth import login_required
from config import COLLECT_IDENTITY, TMP_DIR
from skills import EXPERIENCE, INTENT, SKILL_LABEL, SLEEP

bp = Blueprint("data", __name__, url_prefix="/data")

ALLOWED_SUFFIX = {".xlsx", ".xlsm", ".csv"}


def _actor() -> str:
    from auth import current_user

    user = current_user()
    return user["username"] if user else "system"


# ================================================================== 导入
@bp.route("/import")
@login_required
def import_home():
    return render_template(
        "import.html",
        field_rules=importer.FIELD_RULES,
        batches=models.recent_batches(10),
    )


def _save_upload(file_storage) -> Path:
    suffix = Path(file_storage.filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIX:
        abort(400, "只支持 .xlsx / .xlsm / .csv 文件")
    target = TMP_DIR / f"{uuid.uuid4().hex}{suffix}"
    file_storage.save(target)
    return target


def _safe_tmp(raw: str) -> Path:
    """只允许读 TMP_DIR 里的文件，防止有人构造路径读服务器上的其它文件。"""
    path = (TMP_DIR / Path(raw).name).resolve()
    if TMP_DIR.resolve() not in path.parents or not path.exists():
        abort(400, "上传文件已过期，请重新上传。")
    return path


def _mapping_from_form(form) -> dict:
    mapping = {}
    for key, _label, _kw in importer.FIELD_RULES:
        value = form.get(f"map_{key}")
        if value:
            mapping[key] = value
    return mapping


def _render_preview(tmp: Path, filename: str, mapping: dict | None = None, sheet=None):
    headers, rows, sheets = importer.read_table(tmp, sheet)
    if mapping is None:
        mapping = importer.guess_mapping(headers)
    parsed, warnings = importer.build_members(rows, mapping)
    return render_template(
        "import_preview.html",
        tmp=tmp.name,
        filename=filename,
        headers=headers,
        sheets=sheets,
        sheet=sheet,
        rows_count=len(rows),
        mapping=mapping,
        field_rules=importer.FIELD_RULES,
        preview=parsed[:8],
        parsed_count=len(parsed),
        warnings=warnings[:10],
        skill_label=SKILL_LABEL,
        experience_labels=EXPERIENCE,
        intent_labels=INTENT,
    )


@bp.route("/import/upload", methods=["POST"])
@login_required
def import_upload():
    file_storage = request.files.get("file")
    if not file_storage or not file_storage.filename:
        flash("没有选择文件。", "error")
        return redirect(url_for("data.import_home"))
    tmp = _save_upload(file_storage)
    return _render_preview(tmp, file_storage.filename)


@bp.route("/import/preview", methods=["POST"])
@login_required
def import_preview():
    tmp = _safe_tmp(request.form.get("tmp", ""))
    mapping = _mapping_from_form(request.form)
    return _render_preview(tmp, request.form.get("filename", tmp.name), mapping,
                           request.form.get("sheet") or None)


@bp.route("/import/confirm", methods=["POST"])
@login_required
def import_confirm():
    tmp = _safe_tmp(request.form.get("tmp", ""))
    filename = request.form.get("filename") or tmp.name
    mapping = _mapping_from_form(request.form)
    update_existing = request.form.get("update_existing") == "1"
    if not mapping.get("name"):
        flash("至少要指定「昵称」对应哪一列，否则没法建档案。", "error")
        return _render_preview(tmp, filename, mapping)

    headers, rows, _sheets = importer.read_table(tmp)
    parsed, warnings = importer.build_members(rows, mapping)
    if not parsed:
        flash("没有解析出任何有效成员，检查一下昵称列是不是选错了。", "error")
        return _render_preview(tmp, filename, mapping)

    result = importer.run_import(parsed, filename, mapping, _actor(), update_existing)
    flash(
        f"导入完成：新增 {result['created']} 人，更新 {result['updated']} 人，"
        f"跳过 {result['skipped']} 行。",
        "ok",
    )
    for warning in warnings[:5]:
        flash(warning, "warn")
    return redirect(url_for("members.listing"))


# ================================================================== 导出
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, widths=None):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width


def _send(workbook: Workbook, filename: str):
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _skills_text(skills: dict) -> str:
    if not isinstance(skills, dict):
        return ""
    return "; ".join(f"{SKILL_LABEL.get(k, k)}{v}" for k, v in
                     sorted(skills.items(), key=lambda x: -x[1]))


@bp.route("/export/members.xlsx")
@login_required
def export_members():
    members = models.list_members(status="all")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "成员总表"
    headers = ["会员号", "昵称"]
    if COLLECT_IDENTITY:
        headers.append("真实姓名")
    headers += ["年龄", "经验", "参赛次数", "意向", "求带", "可带新"]
    if COLLECT_IDENTITY:
        headers.append("手机")
    headers += ["可投入小时", "作息", "可语音", "技能", "QQ", "Discord", "邮箱",
                "所在队伍", "队内角色", "状态", "最好成绩", "备注"]
    ws.append(headers)
    for member in members:
        row = [member["code"], member["name"]]
        if COLLECT_IDENTITY:
            row.append(member.get("real_name") or "")
        row += [
            member.get("age") or "",
            EXPERIENCE.get(member["experience"], member["experience"]),
            member["hunts_played"], INTENT.get(member["intent"], member["intent"]),
            "是" if member["wants_mentor"] else "",
            member["mentor_capacity"] or "",
        ]
        if COLLECT_IDENTITY:
            row.append(member.get("phone") or "")
        row += [
            member["hours_available"] or "",
            SLEEP.get(member["sleep_schedule"], member["sleep_schedule"]),
            "是" if member["voice_ok"] else "否",
            _skills_text(member["skills"]),
            member.get("contact_qq") or "", member.get("contact_discord") or "",
            member.get("contact_email") or "",
            member.get("team_name") or "未分配", member.get("team_role") or "",
            member.get("status") or "", member.get("best_result") or "",
            member.get("notes") or "",
        ]
        ws.append(row)
    _style_header(ws)
    return _send(workbook, "成员总表.xlsx")


@bp.route("/export/teams.xlsx")
@login_required
def export_teams():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "队伍名册"
    ws.append(["队伍", "类型", "状态", "人数", "上限", "新手", "新手名额", "队长",
               "会员号", "成员", "经验", "队内角色", "技能"])
    for team in models.list_teams():
        roster = models.team_roster(team["id"])
        if not roster:
            ws.append([team["name"], team["kind"], team["status"], 0, team["capacity"],
                       team["newbie_quota"], "", team.get("captain_name") or ""])
            continue
        for member in roster:
            ws.append([
                team["name"], team["kind"], team["status"], team["size"], team["capacity"],
                team["newbie_count"], team["newbie_quota"], team.get("captain_name") or "",
                member["code"], member["name"],
                EXPERIENCE.get(member["experience"], member["experience"]),
                member.get("team_role") or "",
                _skills_text(member["skills"]),
            ])
    _style_header(ws, [18, 10, 10, 6, 6, 6, 10, 12, 10, 14, 8, 10, 28])
    return _send(workbook, "队伍名册.xlsx")


@bp.route("/export/event/<int:event_id>.xlsx")
@login_required
def export_event(event_id):
    event = models.get_event(event_id)
    if not event:
        abort(404)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "赛事报名"
    ws.append(["队伍", "类型", "人数", "会员号", "成员", "经验", "队内角色", "联系方式"])
    for registration in models.event_registrations(event_id):
        roster = models.team_roster(registration["team_id"])
        for member in roster:
            contact = member.get("contact_discord") or member.get("contact_qq") or ""
            ws.append([
                registration["team_name"], registration["team_kind"], registration["size"],
                member["code"], member["name"],
                EXPERIENCE.get(member["experience"], member["experience"]),
                member.get("team_role") or "", contact,
            ])
    _style_header(ws, [18, 10, 6, 10, 14, 8, 10, 20])
    return _send(workbook, f"{event['name']}_报名表.xlsx")


@bp.route("/export/unassigned.xlsx")
@login_required
def export_unassigned():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "未分配名单"
    ws.append(["会员号", "昵称", "经验", "意向", "求带", "可带新", "技能", "QQ", "Discord", "备注"])
    for member in models.unassigned_members():
        ws.append([
            member["code"], member["name"],
            EXPERIENCE.get(member["experience"], member["experience"]),
            INTENT.get(member["intent"], member["intent"]),
            "是" if member["wants_mentor"] else "",
            member["mentor_capacity"] or "",
            _skills_text(member["skills"]),
            member.get("contact_qq") or "", member.get("contact_discord") or "",
            member.get("notes") or "",
        ])
    _style_header(ws, [10, 14, 8, 8, 6, 8, 28, 14, 18, 24])
    return _send(workbook, "未分配名单.xlsx")


@bp.route("/template.xlsx")
@login_required
def download_template():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "成员导入模板"
    ws.append([label for _key, label, _kw in importer.FIELD_RULES])
    _style_header(ws)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name="成员导入模板.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/export")
@login_required
def export_home():
    return render_template("export.html", events=models.list_events(), batches=models.recent_batches(20))
